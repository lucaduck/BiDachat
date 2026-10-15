import asyncio
import shutil
import subprocess
import sys
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document as WordDocument
from sqlalchemy import delete

from app.core.config import Settings
from app.database.session import Database
from app.models import Chatbot, LlmModel, User
from app.services.document_processing_service import DocumentProcessingService
from app.services.document_service import DocumentService, UploadDocument
from app.services.rag_knowledge_service import RagKnowledgeService


def _docx_with_table() -> bytes:
    document = WordDocument()
    document.add_paragraph("Informe de retención")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Indicador"
    table.cell(0, 1).text = "Valor"
    table.cell(1, 0).text = "RETENCION_BI"
    table.cell(1, 1).text = "87 por ciento"
    content = BytesIO()
    document.save(content)
    return content.getvalue()


@pytest.mark.integration
def test_ingestion_recovers_text_tables_csv_and_images_with_chatbot_isolation(
    tmp_path: Path,
):
    settings = Settings(
        document_storage_path=tmp_path,
        embedding_provider="ollama",
        embedding_model="fixture-embedding",
        ollama_model="fixture-vision",
    )
    if settings.database_url is None:
        pytest.skip("DATABASE_URL is not configured")

    pdf_content = None
    if shutil.which("wkhtmltopdf"):
        source = Path(__file__).parent / "fixtures/table-digital.html"
        pdf = tmp_path / "table-digital.pdf"
        subprocess.run(
            ["wkhtmltopdf", "--quiet", str(source), str(pdf)],
            capture_output=True,
            check=True,
        )
        pdf_content = pdf.read_bytes()

    loop_factory = asyncio.SelectorEventLoop if sys.platform == "win32" else None
    with asyncio.Runner(loop_factory=loop_factory) as runner:
        runner.run(_exercise_ingestion(settings, pdf_content))


async def _exercise_ingestion(settings: Settings, pdf_content: bytes | None) -> None:
    database = Database.from_settings(settings)
    documents = DocumentService(settings)
    processing = DocumentProcessingService(settings)
    knowledge = RagKnowledgeService(settings)
    events: list[str] = []

    async def describe_image(**kwargs) -> str:
        assert kwargs["provider"] == "ollama"
        assert kwargs["model"] == "fixture-vision"
        assert kwargs["image_mime_type"] == "image/png"
        events.append("describe")
        return "La captura muestra VENTAS_VISUALES: 42 unidades."

    async def embed_documents(texts: list[str]) -> list[list[float]]:
        events.append("embed")
        assert texts
        return [[1.0, *([0.0] * 767)] for _ in texts]

    processing.provider.answer = describe_image
    processing.provider.embed_documents = embed_documents

    user_id, model_id = uuid4(), uuid4()
    chatbot_a_id, chatbot_b_id = uuid4(), uuid4()
    try:
        async with database.session() as session:
            session.add_all(
                [
                    User(
                        id=user_id,
                        email=f"ingestion-{user_id}@bidachat.test",
                        password_hash="integration-hash",
                    ),
                    LlmModel(id=model_id, provider="ollama", model="fixture-vision"),
                ]
            )
            await session.flush()
            session.add_all(
                [
                    Chatbot(
                        id=chatbot_a_id,
                        created_by=user_id,
                        configured_llm_model_id=model_id,
                        name="Fixture A",
                    ),
                    Chatbot(
                        id=chatbot_b_id,
                        created_by=user_id,
                        configured_llm_model_id=model_id,
                        name="Fixture B",
                    ),
                ]
            )
            await session.commit()

        uploads = [
            (
                chatbot_a_id,
                UploadDocument(
                    "table.docx",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    _docx_with_table(),
                ),
            ),
            (
                chatbot_a_id,
                UploadDocument(
                    "metrics.csv",
                    "application/csv",
                    b"indicator;value\nVENTAS_CSV;73\n",
                ),
            ),
            (
                chatbot_a_id,
                UploadDocument("chart.png", "image/png", b"\x89PNG\r\n\x1a\nfixture"),
            ),
            (chatbot_b_id, UploadDocument("private.txt", "text/plain", b"SOLO_B: 19")),
        ]
        if pdf_content is not None:
            uploads.append(
                (
                    chatbot_a_id,
                    UploadDocument("table-digital.pdf", "application/pdf", pdf_content),
                )
            )
        uploaded = []
        for chatbot_id, upload in uploads:
            async with database.session() as session:
                document = await documents.upload_document(
                    session, chatbot_id=chatbot_id, upload=upload
                )
                uploaded.append((chatbot_id, document.id))
                await session.commit()
            async with database.session() as session:
                document = await documents.get_document(
                    session, chatbot_id=chatbot_id, document_id=uploaded[-1][1]
                )
                processed = await processing.process(
                    session, document=document, chatbot_id=chatbot_id
                )
                assert processed.status == "ready"

        assert events == ["embed", "embed", "describe", "embed", "embed"] + (
            ["embed"] if pdf_content is not None else []
        )

        async with database.session() as session:
            query_vector = [1.0, *([0.0] * 767)]

            async def retrieve(chatbot_id):
                return await knowledge.retrieve_chunks(
                    session,
                    chatbot_id=chatbot_id,
                    embedding_model=settings.embedding_model,
                    embedding_dimensions=768,
                    query_embedding=query_vector,
                    top_k=50,
                )

            a_text = " ".join(chunk.content for chunk in await retrieve(chatbot_a_id))
            b_text = " ".join(chunk.content for chunk in await retrieve(chatbot_b_id))
            assert "Indicador: RETENCION_BI" in a_text
            assert "value: 73" in a_text
            assert "VENTAS_VISUALES: 42" in a_text
            if pdf_content is not None:
                assert "PDF_TASA_BI" in a_text
            assert "SOLO_B" not in a_text
            assert "SOLO_B: 19" in b_text
            assert "RETENCION_BI" not in b_text

            await documents.associate_documents(
                session, chatbot_id=chatbot_b_id, document_ids=[uploaded[0][1]]
            )
            await session.commit()
            shared_text = " ".join(
                chunk.content for chunk in await retrieve(chatbot_b_id)
            )
            assert "RETENCION_BI" in shared_text
            assert "VENTAS_CSV" not in shared_text
            assert "VENTAS_VISUALES" not in shared_text

            await documents.remove_association(
                session, chatbot_id=chatbot_b_id, document_id=uploaded[0][1]
            )
            await session.commit()
            assert "RETENCION_BI" not in " ".join(
                chunk.content for chunk in await retrieve(chatbot_b_id)
            )
            assert "RETENCION_BI" in " ".join(
                chunk.content for chunk in await retrieve(chatbot_a_id)
            )
    finally:
        async with database.session() as session:
            await session.execute(delete(Chatbot).where(Chatbot.created_by == user_id))
            user = await session.get(User, user_id)
            if user is not None:
                await session.delete(user)
            model = await session.get(LlmModel, model_id)
            if model is not None:
                await session.delete(model)
            await session.commit()
        await database.dispose()
