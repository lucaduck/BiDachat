import asyncio
import shutil
import subprocess
from pathlib import Path

import pytest
from docx import Document as WordDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.core.config import Settings
from app.llm.provider import ProviderConfigurationError
from app.services.document_processing_service import (
    DOCX_MEDIA_TYPE,
    DocumentProcessingService,
)


def test_docx_table_keeps_its_position_and_cell_context(tmp_path):
    document = WordDocument()
    document.add_paragraph("PERIODO ANTERIOR")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Indicador"
    table.cell(0, 1).text = "Valor"
    table.cell(1, 0).text = "RETENCION_BI"
    table.cell(1, 1).text = "87 por ciento"
    document.add_paragraph("PERIODO SIGUIENTE")
    path = tmp_path / "ordered-table.docx"
    document.save(path)

    text = DocumentProcessingService._extract_text(path, DOCX_MEDIA_TYPE)

    assert text.index("PERIODO ANTERIOR") < text.index("RETENCION_BI")
    assert text.index("RETENCION_BI") < text.index("PERIODO SIGUIENTE")
    assert "Indicador" in text and "87 por ciento" in text
    assert "Indicador: RETENCION_BI" in text
    assert "Valor: 87 por ciento" in text


def test_csv_late_row_retains_column_names_when_split(tmp_path):
    path = tmp_path / "metrics.csv"
    rows = ["indicator,value,period"]
    rows.extend(f"metric_{index},{index * 10},2026-Q1" for index in range(150))
    path.write_text("\n".join(rows), encoding="utf-8")

    text = DocumentProcessingService._extract_text(path, "text/csv")
    chunks = RecursiveCharacterTextSplitter(
        chunk_size=1200, chunk_overlap=150
    ).split_text(text)

    assert any("metric_149" in chunk and "indicator" in chunk for chunk in chunks)


def test_csv_alias_and_quoted_semicolon_rows_keep_column_context(tmp_path):
    path = tmp_path / "metrics.csv"
    path.write_text(
        'indicator;value;note\nRETENCION_BI;87;"primer; semestre"\n',
        encoding="utf-8",
    )

    for media_type in ("text/csv", "application/csv", "application/vnd.ms-excel"):
        text = DocumentProcessingService._extract_text(path, media_type)
        assert "indicator: RETENCION_BI" in text
        assert "value: 87" in text
        assert "note: primer; semestre" in text


def test_digital_pdf_table_exposes_known_cells(tmp_path):
    if shutil.which("wkhtmltopdf") is None:
        pytest.skip("wkhtmltopdf is unavailable")
    source = Path(__file__).parent / "fixtures/table-digital.html"
    pdf = tmp_path / "table-digital.pdf"
    subprocess.run(
        ["wkhtmltopdf", "--quiet", str(source), str(pdf)],
        capture_output=True,
        check=True,
    )

    text = DocumentProcessingService._extract_text(pdf, "application/pdf")

    assert "PDF_TASA_BI" in text
    assert "91 por" in text
    assert "ciento" in text


@pytest.mark.parametrize("provider", ["ollama", "gemini", "openai", "openrouter"])
def test_image_source_uses_explicit_vision_profile(tmp_path, provider):
    path = tmp_path / "chart.png"
    path.write_bytes(b"fixture image")
    processing = DocumentProcessingService(
        Settings(
            document_storage_path=tmp_path,
            ollama_model="text-only-local",
            vision_provider=provider,
            vision_model="vision-fixture",
        )
    )
    calls = []

    async def describe(**kwargs):
        calls.append(kwargs)
        return "Gráfico con 42 ventas"

    processing.provider.answer = describe

    assert asyncio.run(processing._extract_content(path, "image/png")) == (
        "Gráfico con 42 ventas"
    )
    assert calls[0]["provider"] == provider
    assert calls[0]["model"] == "vision-fixture"
    assert calls[0]["image"] == b"fixture image"


def test_image_source_without_vision_model_fails_explicitly(tmp_path):
    path = tmp_path / "chart.png"
    path.write_bytes(b"fixture image")
    processing = DocumentProcessingService(
        Settings(
            document_storage_path=tmp_path,
            ollama_model="",
            vision_provider="openai",
            vision_model="",
        )
    )

    with pytest.raises(ProviderConfigurationError, match="VISION_MODEL"):
        asyncio.run(processing._extract_content(path, "image/png"))
